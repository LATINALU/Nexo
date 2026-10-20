"""
Nexo - Documentos adjuntos (RAG)
================================
El usuario sube archivos; se extrae el texto, se divide en fragmentos y en cada
consulta se recuperan los más relevantes (BM25) para dárselos a los agentes.

Almacenamiento: PostgreSQL si está disponible; si no, memoria del proceso (24 h).
Acceso: el id del documento es un UUID aleatorio (capacidad). Si el documento
pertenece a un usuario registrado, además se exige ser ese usuario.
"""

from __future__ import annotations

import io
import json
import logging
import time
import uuid
from dataclasses import dataclass, field
from typing import Dict, List, Optional

from app.services.text_utils import BM25, chunk_text

logger = logging.getLogger("nexo.documents")

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_DOCUMENT_CHARS = 300_000
MAX_DOCUMENTS_PER_RUN = 5
MAX_CONTEXT_CHARS = 7000
TOP_CHUNKS = 6
_MEMORY_TTL_S = 24 * 3600
_MEMORY_MAX_DOCS = 200

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm", ".log")


class DocumentError(ValueError):
    """Error presentable al usuario."""


@dataclass
class Document:
    id: str
    name: str
    mime: str
    size: int
    owner_id: Optional[str]
    chunks: List[str]
    created_at: float = field(default_factory=time.time)

    @property
    def char_count(self) -> int:
        return sum(len(c) for c in self.chunks)

    def public(self) -> Dict:
        return {"id": self.id, "name": self.name, "mime": self.mime, "size": self.size,
                "chunks": len(self.chunks), "chars": self.char_count}


# ----------------------------------------------------------------- extracción

def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if not name.endswith(SUPPORTED_EXTENSIONS):
        raise DocumentError(f"Formato no soportado. Usa: {', '.join(SUPPORTED_EXTENSIONS)}")
    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader
        except ImportError:
            raise DocumentError("Soporte PDF no instalado (pypdf)")
        try:
            reader = PdfReader(io.BytesIO(data))
            text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as e:
            raise DocumentError(f"No se pudo leer el PDF: {e}"[:200])
        if not text.strip():
            raise DocumentError("El PDF no tiene texto extraíble (¿es un escaneo?)")
        return text
    if name.endswith(".docx"):
        try:
            import docx
        except ImportError:
            raise DocumentError("Soporte DOCX no instalado (python-docx)")
        try:
            document = docx.Document(io.BytesIO(data))
        except Exception as e:
            raise DocumentError(f"No se pudo leer el DOCX: {e}"[:200])
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n\n".join(parts)
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if name.endswith((".html", ".htm")):
        import re
        text = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", text, flags=re.S | re.I)
        text = re.sub(r"<[^>]+>", " ", text)
    if name.endswith(".json"):
        try:
            text = json.dumps(json.loads(text), ensure_ascii=False, indent=1)
        except ValueError:
            pass
    return text


def build_document(filename: str, mime: str, data: bytes, owner_id: Optional[str]) -> Document:
    if len(data) > MAX_FILE_BYTES:
        raise DocumentError(f"El archivo supera {MAX_FILE_BYTES // (1024 * 1024)} MB")
    if not data:
        raise DocumentError("El archivo está vacío")
    text = extract_text(filename, data)[:MAX_DOCUMENT_CHARS]
    chunks = chunk_text(text)
    if not chunks:
        raise DocumentError("No se encontró texto en el archivo")
    safe_name = (filename.replace("\\", "/").split("/")[-1] or "documento")[:200]
    return Document(id=str(uuid.uuid4()), name=safe_name, mime=(mime or "")[:100], size=len(data),
                    owner_id=owner_id, chunks=chunks)


# ---------------------------------------------------------------- almacenamiento

class DocumentStore:

    def __init__(self):
        self._memory: Dict[str, Document] = {}

    @property
    def _pool(self):
        from app.pg_database import pg_db
        return pg_db._pool

    def _purge_memory(self) -> None:
        now = time.time()
        for doc_id in [d.id for d in self._memory.values() if now - d.created_at > _MEMORY_TTL_S]:
            self._memory.pop(doc_id, None)
        while len(self._memory) > _MEMORY_MAX_DOCS:
            oldest = min(self._memory.values(), key=lambda d: d.created_at)
            self._memory.pop(oldest.id, None)

    async def save(self, doc: Document) -> None:
        if self._pool:
            async with self._pool.acquire() as conn, conn.transaction():
                await conn.execute("""
                    INSERT INTO public.nexo_documents (id, owner_id, name, mime, size, char_count, chunk_count)
                    VALUES ($1, $2, $3, $4, $5, $6, $7)
                """, uuid.UUID(doc.id), uuid.UUID(doc.owner_id) if doc.owner_id else None, doc.name, doc.mime,
                    doc.size, doc.char_count, len(doc.chunks))
                await conn.executemany(
                    "INSERT INTO public.nexo_document_chunks (document_id, ord, content) VALUES ($1, $2, $3)",
                    [(uuid.UUID(doc.id), i, c) for i, c in enumerate(doc.chunks)])
            return
        self._memory[doc.id] = doc
        self._purge_memory()

    async def get(self, doc_id: str, user_id: Optional[str]) -> Optional[Document]:
        try:
            uid = uuid.UUID(doc_id)
        except (ValueError, TypeError):
            return None
        if self._pool:
            async with self._pool.acquire() as conn:
                row = await conn.fetchrow("SELECT * FROM public.nexo_documents WHERE id = $1", uid)
                if not row:
                    return None
                owner = str(row["owner_id"]) if row["owner_id"] else None
                if owner and owner != user_id:
                    return None
                chunks = await conn.fetch(
                    "SELECT content FROM public.nexo_document_chunks WHERE document_id = $1 ORDER BY ord", uid)
            return Document(id=doc_id, name=row["name"], mime=row["mime"] or "", size=row["size"],
                            owner_id=owner, chunks=[c["content"] for c in chunks],
                            created_at=row["created_at"].timestamp())
        doc = self._memory.get(str(uid))
        if not doc or (doc.owner_id and doc.owner_id != user_id):
            return None
        return doc

    async def list_for_user(self, user_id: str) -> List[Dict]:
        if self._pool:
            async with self._pool.acquire() as conn:
                rows = await conn.fetch("""
                    SELECT id, name, mime, size, chunk_count, char_count FROM public.nexo_documents
                    WHERE owner_id = $1 ORDER BY created_at DESC LIMIT 200
                """, uuid.UUID(user_id))
            return [{"id": str(r["id"]), "name": r["name"], "mime": r["mime"], "size": r["size"],
                     "chunks": r["chunk_count"], "chars": r["char_count"]} for r in rows]
        return [d.public() for d in self._memory.values() if d.owner_id == user_id]

    async def delete(self, doc_id: str, user_id: Optional[str]) -> bool:
        doc = await self.get(doc_id, user_id)
        if not doc:
            return False
        if self._pool:
            async with self._pool.acquire() as conn:
                await conn.execute("DELETE FROM public.nexo_documents WHERE id = $1", uuid.UUID(doc_id))
        else:
            self._memory.pop(doc_id, None)
        return True


document_store = DocumentStore()


# ------------------------------------------------------------------ retrieval

async def load_documents(doc_ids: List[str], user_id: Optional[str]) -> List[Document]:
    docs = []
    for doc_id in list(dict.fromkeys(doc_ids))[:MAX_DOCUMENTS_PER_RUN]:
        doc = await document_store.get(str(doc_id), user_id)
        if doc:
            docs.append(doc)
    return docs


def retrieve(query: str, docs: List[Document], k: int = TOP_CHUNKS,
             max_chars: int = MAX_CONTEXT_CHARS) -> List[Dict]:
    """
    Fragmentos más relevantes para la consulta, etiquetados [D1], [D2]… por documento.
    Si ningún fragmento coincide léxicamente (p.ej. "resume este archivo"),
    se usa el inicio de cada documento.
    """
    if not docs:
        return []
    flat = [(d_idx, c_idx, chunk) for d_idx, d in enumerate(docs) for c_idx, chunk in enumerate(d.chunks)]
    ranked = BM25([c for _, _, c in flat]).top(query, k)
    picks = [flat[i] for i, _ in ranked]
    if not picks:
        per_doc = max(1, k // len(docs))
        picks = [(d_idx, c_idx, d.chunks[c_idx]) for d_idx, d in enumerate(docs)
                 for c_idx in range(min(per_doc, len(d.chunks)))]
    picks.sort(key=lambda p: (p[0], p[1]))  # orden de lectura

    results, used = [], 0
    for d_idx, c_idx, chunk in picks:
        if used >= max_chars:
            break
        text = chunk[: max_chars - used]
        used += len(text)
        results.append({"label": f"D{d_idx + 1}", "document_id": docs[d_idx].id, "name": docs[d_idx].name,
                        "chunk": c_idx + 1, "content": text})
    return results
